from __future__ import annotations

import argparse
import hashlib
import importlib.util
import re
import shutil
import tempfile
from copy import deepcopy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "deliverables" / "论文内容底稿.md"
DEFAULT_TEMPLATE = Path(r"C:\Users\Administrator\Desktop\毕业设计论文模板.docx")
DEFAULT_OUTPUT = ROOT / "deliverables" / "毕业论文.docx"
FIG_DIR = ROOT / "artifacts" / "docx-qa" / "figures"
SCREENSHOT_DIR = ROOT / "artifacts" / "e2e-vue-final"

EXPECTED_TEMPLATE_SHA256 = "8EA289BED8664F323722129CCFEA82CD3F1AF304E509956509DC7881732EB361"
TITLE = "基于SpringBoot与WebSocket的多人实时协同编辑系统的设计与实现"
TITLE_LINE_1 = "基于SpringBoot与WebSocket的"
TITLE_LINE_2 = "多人实时协同编辑系统的设计与实现"
ABSTRACT_TITLE_LINES = (TITLE_LINE_1, TITLE_LINE_2)
ENGLISH_ABSTRACT_TITLE_LINES = (
    "Real-time Collaborative Editing System",
    "Based on Spring Boot and WebSocket",
)

FONT_BODY = "宋体"
FONT_HEADING = "黑体"
FONT_LATIN = "Times New Roman"
FONT_MATH = "Cambria Math"
BLACK = "000000"
GRID = "B7B7B7"

FIGURES = {
    "1-1": ("研究技术路线", "route"),
    "3-1": ("系统总体架构", "architecture"),
    "3-2": ("系统功能模块", "modules"),
    "3-3": ("LWW 字段合并流程", "lww"),
    "3-4": ("WebSocket 连接状态", "ws_state"),
    "3-5": ("断线恢复时序", "reconnect"),
    "3-6": ("操作提交时序", "commit"),
    "3-7": ("数据库逻辑 ER 关系（当前未声明 SQL FOREIGN KEY）", "er"),
    "4-1": ("协同文档双端同步界面", "03-document-owner-editor-sync.png"),
    "4-2": ("协同白板卡片同步界面", "06-board-synced.png"),
    "4-3": ("协同看板任务卡跨列后的同步结果", "07-kanban-drag-synced.png"),
    "4-4": ("工作台空间与成员管理界面", "02-owner-members.png"),
    "4-5": ("版本历史恢复结果", "04-history-restored.png"),
    "4-6": ("正文搜索与协同统计界面", "09-search-result.png"),
    "5-1": ("Playwright 三角色权限验证", ("02-owner-members.png", "05-viewer-document-readonly.png")),
    "5-2": ("WebSocket 中断 30003 ms 后自动重连并补齐", "08-editor-reconnected-after-30s.png"),
    "5-3": ("不同客户端规模下的操作确认延迟", "latency"),
    "5-4": ("不同客户端规模下的吞吐量", "throughput"),
}

TABLE_CAPTIONS = [
    ("2-1", "核心权限矩阵"),
    ("3-1", "三类应用的统一模型投影"),
    ("5-1", "测试环境"),
    ("5-2", "服务端集成测试覆盖"),
    ("5-3", "Maven clean 测试结果"),
    ("5-4", "并发负载延迟与吞吐结果"),
    ("5-5", "持久化和版本完整性结果"),
]

SPECIAL_HEADINGS = {
    "结论": "结    论",
    "参考文献": "参 考 文 献",
    "致谢": "致    谢",
}

SUBSCRIPT = str.maketrans(
    {
        "0": "₀",
        "1": "₁",
        "2": "₂",
        "3": "₃",
        "4": "₄",
        "5": "₅",
        "6": "₆",
        "7": "₇",
        "8": "₈",
        "9": "₉",
        "a": "ₐ",
        "e": "ₑ",
        "h": "ₕ",
        "i": "ᵢ",
        "j": "ⱼ",
        "k": "ₖ",
        "l": "ₗ",
        "m": "ₘ",
        "n": "ₙ",
        "o": "ₒ",
        "p": "ₚ",
        "r": "ᵣ",
        "s": "ₛ",
        "t": "ₜ",
        "u": "ᵤ",
        "v": "ᵥ",
        "x": "ₓ",
        "β": "ᵦ",
    }
)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def load_figure_generator():
    script = ROOT / "scripts" / "build-thesis-docx.py"
    spec = importlib.util.spec_from_file_location("legacy_thesis_figures", script)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"cannot import figure generator: {script}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def set_run_font(run, *, east=FONT_BODY, latin=FONT_LATIN, size=12, bold=None, italic=None) -> None:
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), east)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(BLACK)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def clear_paragraph(paragraph: Paragraph) -> None:
    p_pr = paragraph._p.pPr
    for child in list(paragraph._p):
        if child is not p_pr:
            paragraph._p.remove(child)


def set_paragraph_text(
    paragraph: Paragraph,
    text: str,
    *,
    east=FONT_BODY,
    latin=FONT_LATIN,
    size=12,
    bold=None,
    italic=None,
) -> None:
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, east=east, latin=latin, size=size, bold=bold, italic=italic)


def set_multiline_title(paragraph: Paragraph) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(TITLE_LINE_1)
    set_run_font(run, east=FONT_HEADING, latin=FONT_LATIN, size=26)
    run.add_break()
    run = paragraph.add_run(TITLE_LINE_2)
    set_run_font(run, east=FONT_HEADING, latin=FONT_LATIN, size=26)


def replace_exact_text_nodes(element, replacements: dict[str, str]) -> None:
    for node in element.xpath(".//w:t"):
        if node.text in replacements:
            node.text = replacements[node.text]


def replace_direct_text(paragraph: Paragraph, text: str) -> None:
    nodes = paragraph._p.xpath("./w:r/w:t")
    if not nodes:
        run = paragraph.add_run(text)
        set_run_font(run, size=16)
        return
    nodes[0].text = text
    for node in nodes[1:]:
        node.text = ""


def replace_direct_text_lines(
    paragraph: Paragraph,
    lines: tuple[str, ...],
    *,
    east: str,
    latin: str,
    size: float,
    line_spacing: float,
) -> None:
    direct_runs = paragraph._p.xpath("./w:r[w:t]")
    if direct_runs:
        primary = direct_runs[0]
        for extra in direct_runs[1:]:
            paragraph._p.remove(extra)
        r_pr = primary.rPr
        for child in list(primary):
            if child is not r_pr:
                primary.remove(child)
        run = Run(primary, paragraph)
    else:
        run = paragraph.add_run()

    set_run_font(run, east=east, latin=latin, size=size, bold=False)
    for index, line in enumerate(lines):
        if index:
            run.add_break()
        run.add_text(line)

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(line_spacing)


def set_metadata_vertical_position(
    paragraph: Paragraph,
    *,
    offset_emu: int,
    margin_top_pt: float,
) -> None:
    for node in paragraph._p.xpath(".//wp:anchor/wp:positionV/wp:posOffset"):
        node.text = str(offset_emu)
    for shape in paragraph._p.xpath('.//*[local-name()="shape"]'):
        style = shape.get("style")
        if not style or "margin-top:" not in style:
            continue
        shape.set(
            "style",
            re.sub(r"margin-top:[^;]+", f"margin-top:{margin_top_pt:g}pt", style),
        )


def replace_cover_identity(paragraph: Paragraph) -> None:
    for node in paragraph._p.xpath(".//w:t"):
        value = node.text or ""
        if "电气与信息工程学院" in value:
            node.text = re.sub(r"电气与信息工程学院", "", value)
        elif "计算机科学与技术" in value:
            node.text = value.replace("计算机科学与技术", "软件工程一班")
        elif value.strip() == "XXX" and value.startswith("  ") and not value.startswith("        "):
            node.text = value.replace("XXX", "柯佳伟")
        elif "20230000000" in value:
            node.text = value.replace("20230000000", "20235102034")
        elif value.strip() == "XXX" and value.startswith("        "):
            node.text = value.replace("XXX", "   ")


def expand_cover_identity_box(paragraph: Paragraph) -> None:
    height_emu = "2047875"  # 161.25 pt; keeps the reviewer row inside the text box.
    for node in paragraph._p.xpath(".//wp:extent | .//a:ext"):
        if node.get("cy") == "1895475":
            node.set("cy", height_emu)
    for node in paragraph._p.iter("{urn:schemas-microsoft-com:vml}shape"):
        style = node.get("style", "")
        node.set("style", re.sub(r"height:[0-9.]+pt", "height:161.25pt", style))


def replace_abstract_metadata(paragraph: Paragraph) -> None:
    major_tail_pending = False
    for node in paragraph._p.xpath(".//w:t"):
        value = node.text or ""
        if value == "XX":
            node.text = "柯佳伟"
        elif value == "计算机科学与":
            node.text = "软件工程一班"
            major_tail_pending = True
        elif value == "技术" and major_tail_pending:
            node.text = ""
            major_tail_pending = False
        elif value == "XXX":
            node.text = "揣小龙"
    replace_direct_text_lines(
        paragraph,
        ABSTRACT_TITLE_LINES,
        east=FONT_HEADING,
        latin=FONT_LATIN,
        size=15,
        line_spacing=20,
    )


def replace_english_metadata(paragraph: Paragraph) -> None:
    replacements = ["Ke Jiawei", "Software Engineering", "Chuai Xiaolong"] * 2
    index = 0
    for node in paragraph._p.xpath(".//w:t"):
        if node.text == "Major:":
            node.text = "Major: "
            node.set(qn("xml:space"), "preserve")
        elif node.text == "XXX" and index < len(replacements):
            node.text = replacements[index]
            index += 1
    replace_direct_text_lines(
        paragraph,
        ENGLISH_ABSTRACT_TITLE_LINES,
        east=FONT_LATIN,
        latin=FONT_LATIN,
        size=15,
        line_spacing=20,
    )
    set_metadata_vertical_position(
        paragraph,
        offset_emu=556260,
        margin_top_pt=43.8,
    )


def configure_body_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT_BODY)
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    for level, size in ((1, 16), (2, 14), (3, 12)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT_LATIN
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT_HEADING)
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True
        if level == 1:
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = Pt(11)
        else:
            style.paragraph_format.space_before = Pt(2.5)
            style.paragraph_format.space_after = Pt(0)


def add_body_paragraph(doc: Document, text: str, *, first_indent=True, size=12) -> Paragraph:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.widow_control = True
    paragraph.paragraph_format.first_line_indent = Pt(24) if first_indent else None
    add_inline_content(paragraph, text, size=size)
    return paragraph


def add_plain_run(paragraph: Paragraph, text: str, *, size=12, bold=False, italic=False) -> None:
    if not text:
        return
    clean_text = text.replace("`", "")
    clean_text = re.sub(r"(?<=\d) (?=(?:ms|s|px|MiB|KiB)\b)", "\u00a0", clean_text)
    clean_text = re.sub(r"约 (?=\d+\u00a0(?:ms|s)\b)", "约\u00a0", clean_text)
    run = paragraph.add_run(clean_text)
    set_run_font(run, size=size, bold=bold, italic=italic)


def latex_to_unicode(source: str) -> str:
    text = source.strip()
    text = text.replace("。", "").replace("，", "")
    if "\\begin{cases}" in text:
        return "merge((vₒ, Lₒ), (vₙ, Lₙ)) = { (vₙ, Lₙ), Lₙ > Lₒ; (vₒ, Lₒ), Lₙ ≤ Lₒ }"
    text = text.replace("\\{", "{").replace("\\}", "}")
    text = text.replace("\\ldots", "…")
    text = text.replace("\\Leftrightarrow", "⇔")
    text = text.replace("\\lor", "∨")
    text = text.replace("\\land", "∧")
    text = text.replace("\\le", "≤")
    text = text.replace("\\max", "max")
    text = text.replace("\\,", " ")

    def sub_braced(match: re.Match[str]) -> str:
        return match.group(1) + match.group(2).translate(SUBSCRIPT)

    text = re.sub(r"([A-Za-zβ])_\{([A-Za-z0-9β]+)\}", sub_braced, text)
    text = re.sub(r"([A-Za-zβ])_([A-Za-z0-9β])", sub_braced, text)
    text = text.replace("<ₗₑₓ", "<ₗₑₓ")
    text = text.replace("&", "")
    text = re.sub(r"\\begin\{[^}]+\}|\\end\{[^}]+\}|\\\\", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


MATH_SUBSCRIPT_PATTERN = re.compile(
    r"(?P<base>[A-Za-z\u03b2<])_(?:\{(?P<braced>[A-Za-z0-9\u03b2]+)\}|(?P<single>[A-Za-z0-9\u03b2]))"
)
MATH_CASES_PATTERN = re.compile(
    r"^(?P<prefix>.*?)\\begin\{cases\}(?P<body>.*?)\\end\{cases\}[\u3002\uff0c]?$",
    re.DOTALL,
)


def normalize_math_text(source: str) -> str:
    text = source.strip().replace("\u3002", "").replace("\uff0c", "")
    for latex, replacement in (
        ("\\{", "{"),
        ("\\}", "}"),
        ("\\ldots", "\u2026"),
        ("\\Leftrightarrow", "\u21d4"),
        ("\\lor", "\u2228"),
        ("\\land", "\u2227"),
        ("\\le", "\u2264"),
        ("\\max", "max"),
        ("\\,", " "),
    ):
        text = text.replace(latex, replacement)
    return re.sub(r"\s+", " ", text).strip()


def append_math_run(parent, text: str) -> None:
    if not text:
        return
    math_run = OxmlElement("m:r")
    r_pr = OxmlElement("m:rPr")
    sty = OxmlElement("m:sty")
    sty.set(qn("m:val"), "p")
    r_pr.append(sty)
    math_run.append(r_pr)
    math_text = OxmlElement("m:t")
    math_text.text = text
    math_run.append(math_text)
    parent.append(math_run)


def append_math_expression(parent, source: str) -> None:
    text = normalize_math_text(source)
    cursor = 0
    for match in MATH_SUBSCRIPT_PATTERN.finditer(text):
        append_math_run(parent, text[cursor : match.start()])

        subscript = OxmlElement("m:sSub")
        base = OxmlElement("m:e")
        append_math_run(base, match.group("base"))
        subscript.append(base)

        sub = OxmlElement("m:sub")
        append_math_run(sub, match.group("braced") or match.group("single"))
        subscript.append(sub)
        parent.append(subscript)
        cursor = match.end()

    append_math_run(parent, text[cursor:])


def append_cases(parent, prefix: str, body: str) -> None:
    append_math_expression(parent, prefix)

    delimiter = OxmlElement("m:d")
    delimiter_properties = OxmlElement("m:dPr")
    begin_character = OxmlElement("m:begChr")
    begin_character.set(qn("m:val"), "{")
    delimiter_properties.append(begin_character)
    end_character = OxmlElement("m:endChr")
    end_character.set(qn("m:val"), "")
    delimiter_properties.append(end_character)
    grow = OxmlElement("m:grow")
    grow.set(qn("m:val"), "1")
    delimiter_properties.append(grow)
    delimiter.append(delimiter_properties)

    delimiter_expression = OxmlElement("m:e")
    matrix = OxmlElement("m:m")
    matrix_properties = OxmlElement("m:mPr")
    columns = OxmlElement("m:mcs")
    column = OxmlElement("m:mc")
    column_properties = OxmlElement("m:mcPr")
    column_count = OxmlElement("m:count")
    column_count.set(qn("m:val"), "2")
    column_properties.append(column_count)
    column_justification = OxmlElement("m:mcJc")
    column_justification.set(qn("m:val"), "left")
    column_properties.append(column_justification)
    column.append(column_properties)
    columns.append(column)
    matrix_properties.append(columns)
    matrix.append(matrix_properties)

    rows = [row.strip() for row in re.split(r"\\\\", body) if row.strip()]
    for row_text in rows:
        result, separator, condition = row_text.partition("&")
        row = OxmlElement("m:mr")
        for cell_text in (result, condition if separator else ""):
            cell = OxmlElement("m:e")
            append_math_expression(cell, cell_text)
            row.append(cell)
        matrix.append(row)

    delimiter_expression.append(matrix)
    delimiter.append(delimiter_expression)
    parent.append(delimiter)


def make_omath(source: str):
    o_math = OxmlElement("m:oMath")
    cases = MATH_CASES_PATTERN.match(source.strip())
    if cases:
        append_cases(o_math, cases.group("prefix"), cases.group("body"))
    else:
        append_math_expression(o_math, source)
    return o_math


def add_inline_math(paragraph: Paragraph, latex: str) -> None:
    paragraph._p.append(make_omath(latex))


def add_display_math(doc: Document, latex_lines: list[str]) -> None:
    latex = " ".join(line.strip() for line in latex_lines if line.strip() not in ("\\[", "\\]"))
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_together = True
    math_para = OxmlElement("m:oMathPara")
    math_para_pr = OxmlElement("m:oMathParaPr")
    justification = OxmlElement("m:jc")
    justification.set(qn("m:val"), "center")
    math_para_pr.append(justification)
    math_para.append(math_para_pr)
    math_para.append(make_omath(latex))
    paragraph._p.append(math_para)


INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`.+?`|\\\(.+?\\\))")


def add_inline_content(paragraph: Paragraph, text: str, *, size=12) -> None:
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        add_plain_run(paragraph, text[cursor : match.start()], size=size)
        token = match.group(0)
        if token.startswith("**"):
            add_plain_run(paragraph, token[2:-2], size=size, bold=True)
        elif token.startswith("`"):
            add_plain_run(paragraph, token[1:-1], size=size)
        else:
            add_inline_math(paragraph, token[2:-2])
        cursor = match.end()
    add_plain_run(paragraph, text[cursor:], size=size)


def set_cell_border(cell, *, color=GRID, size="6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def hide_table_borders(table) -> None:
    table_pr = table._tbl.tblPr
    borders = table_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def set_cell_margins(cell, *, top=70, start=75, bottom=70, end=75) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_row_together(row, *, repeat=False) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)
    if repeat:
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_table_separator(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def clean_markdown_cell(value: str) -> str:
    value = re.sub(r"\*\*(.+?)\*\*", r"\1", value)
    value = re.sub(r"`(.+?)`", r"\1", value)
    value = re.sub(r"\\\((.+?)\\\)", lambda m: latex_to_unicode(m.group(1)), value)
    return value.replace("`", "")


def add_table_caption(doc: Document, number: str, title: str) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(f"表{number} {title}")
    set_run_font(run, size=10.5)


def add_markdown_table(doc: Document, rows: list[list[str]], table_index: int) -> None:
    if table_index >= len(TABLE_CAPTIONS):
        raise RuntimeError(f"unexpected table index {table_index}")
    number, title = TABLE_CAPTIONS[table_index]
    add_table_caption(doc, number, title)
    columns = max(len(row) for row in rows)
    normalized = [row + [""] * (columns - len(row)) for row in rows]
    table = doc.add_table(rows=len(normalized), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    total_width_cm = 15.8
    max_lengths = [max(2, max(len(clean_markdown_cell(row[col])) for row in normalized)) for col in range(columns)]
    weights = [min(max(length, 4), 28) for length in max_lengths]
    weight_sum = sum(weights)
    widths = [total_width_cm * weight / weight_sum for weight in weights]
    font_size = 8.5 if columns >= 9 else 9 if columns >= 7 else 9.5 if columns >= 5 else 10.5

    for row_index, row_values in enumerate(normalized):
        row = table.rows[row_index]
        keep_row_together(row, repeat=row_index == 0)
        for column_index, value in enumerate(row_values):
            cell = row.cells[column_index]
            cell.width = Cm(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.style = doc.styles["Table Text"] if "Table Text" in doc.styles else doc.styles["Normal"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 or len(value) <= 22 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = None
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            add_plain_run(paragraph, clean_markdown_cell(value), size=font_size, bold=row_index == 0)


def add_figure_caption(doc: Document, number: str, title: str) -> None:
    style = "图名中文" if "图名中文" in doc.styles else "Normal"
    paragraph = doc.add_paragraph(style=style)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run(f"图{number} {title}")
    set_run_font(run, size=10.5)


def set_picture_alt(run, *, title: str, description: str) -> None:
    doc_pr_nodes = run._r.xpath(".//wp:docPr")
    if not doc_pr_nodes:
        return
    doc_pr = doc_pr_nodes[0]
    doc_pr.set("name", title)
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_image(doc: Document, path: Path, number: str, title: str, *, max_width_cm=15.2, max_height_cm=10.2) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    with Image.open(path) as image:
        aspect = image.width / image.height
    width_cm = max_width_cm
    height_cm = width_cm / aspect
    if height_cm > max_height_cm:
        height_cm = max_height_cm
        width_cm = height_cm * aspect
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    run.add_picture(str(path), width=Cm(width_cm), height=Cm(height_cm))
    set_picture_alt(run, title=f"图{number} {title}", description=title)
    add_figure_caption(doc, number, title)


def add_cropped_image(
    doc: Document,
    path: Path,
    number: str,
    title: str,
    *,
    bottom_ratio: float,
) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    cropped_path = FIG_DIR / f"figure-{number}-cropped.png"
    with Image.open(path) as image:
        crop_bottom = max(1, min(image.height, round(image.height * bottom_ratio)))
        image.crop((0, 0, image.width, crop_bottom)).save(cropped_path)
    add_image(doc, cropped_path, number, title)


def add_two_images(doc: Document, paths: tuple[Path, Path], number: str, title: str) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hide_table_borders(table)
    keep_row_together(table.rows[0])
    for index, (cell, path) in enumerate(zip(table.rows[0].cells, paths, strict=True), start=1):
        if not path.exists():
            raise FileNotFoundError(path)
        cell.width = Cm(7.7)
        set_cell_margins(cell, top=50, start=50, bottom=50, end=50)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = None
        with Image.open(path) as image:
            aspect = image.width / image.height
        width_cm = 7.35
        height_cm = min(5.0, width_cm / aspect)
        run = paragraph.add_run()
        run.add_picture(str(path), width=Cm(width_cm), height=Cm(height_cm))
        set_picture_alt(run, title=f"图{number}-{index}", description=title)
    add_figure_caption(doc, number, title)


def add_figure(doc: Document, number: str, generated: dict[str, Path]) -> None:
    if number not in FIGURES:
        raise RuntimeError(f"figure marker has no mapping: {number}")
    title, source = FIGURES[number]
    if isinstance(source, tuple):
        add_two_images(doc, tuple(SCREENSHOT_DIR / name for name in source), number, title)
    elif number == "4-1":
        add_cropped_image(doc, SCREENSHOT_DIR / source, number, title, bottom_ratio=0.5)
    elif source.endswith(".png"):
        add_image(doc, SCREENSHOT_DIR / source, number, title)
    else:
        add_image(doc, generated[source], number, title, max_height_cm=9.2)


def add_heading(doc: Document, source_text: str, level: int, *, first_heading=False) -> None:
    display_text = SPECIAL_HEADINGS.get(source_text, source_text)
    if source_text in SPECIAL_HEADINGS:
        style = "样式 参考文献标题 + 段后: 1 行" if "样式 参考文献标题 + 段后: 1 行" in doc.styles else "Heading 1"
        paragraph = doc.add_paragraph(style=style)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.page_break_before = not first_heading
        run = paragraph.add_run(display_text)
        set_run_font(run, east=FONT_HEADING, size=16)
        return
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.page_break_before = level == 1 and not first_heading
    paragraph.paragraph_format.keep_with_next = True
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(display_text)
    set_run_font(run, east=FONT_HEADING, size={1: 16, 2: 14, 3: 12}[level])


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph(style="HTML Preformatted" if "HTML Preformatted" in doc.styles else "Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.left_indent = Cm(0.74)
    paragraph.paragraph_format.right_indent = Cm(0.4)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_together = True
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        set_run_font(run, east="微软雅黑", latin="Consolas", size=9.5)
        if index < len(lines) - 1:
            run.add_break()


def parse_body(doc: Document, lines: list[str], generated: dict[str, Path]) -> None:
    index = 0
    table_index = 0
    first_heading = True
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            add_heading(doc, heading.group(2).strip(), level, first_heading=first_heading)
            first_heading = False
            index += 1
            continue
        figure = re.match(r"^\[建议插图(\d+-\d+)：.*\]$", line)
        if figure:
            add_figure(doc, figure.group(1), generated)
            index += 1
            continue
        if line.startswith("```"):
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index].rstrip())
                index += 1
            add_code_block(doc, code_lines)
            index += 1
            continue
        if line == "\\[":
            equation_lines = [line]
            index += 1
            while index < len(lines):
                equation_lines.append(lines[index])
                if lines[index].strip() == "\\]":
                    index += 1
                    break
                index += 1
            add_display_math(doc, equation_lines)
            continue
        if line.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            rows = [split_table_row(line)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_table_row(lines[index]))
                index += 1
            add_markdown_table(doc, rows, table_index)
            table_index += 1
            continue
        if re.match(r"^-\s+", line):
            while index < len(lines) and re.match(r"^-\s+", lines[index].strip()):
                value = re.sub(r"^-\s+", "", lines[index].strip())
                paragraph = doc.add_paragraph(style="Normal")
                paragraph.paragraph_format.left_indent = Cm(0.74)
                paragraph.paragraph_format.first_line_indent = Cm(-0.32)
                paragraph.paragraph_format.line_spacing = 1.25
                add_plain_run(paragraph, "• ", size=12)
                add_inline_content(paragraph, value)
                index += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            item_number = 1
            while index < len(lines) and re.match(r"^\d+\.\s+", lines[index].strip()):
                value = re.sub(r"^\d+\.\s+", "", lines[index].strip())
                paragraph = doc.add_paragraph(style="Normal")
                paragraph.paragraph_format.left_indent = Cm(0.74)
                paragraph.paragraph_format.first_line_indent = Cm(-0.5)
                paragraph.paragraph_format.line_spacing = 1.25
                add_plain_run(paragraph, f"{item_number}. ", size=12)
                add_inline_content(paragraph, value)
                item_number += 1
                index += 1
            continue
        if line.startswith(">"):
            paragraph = doc.add_paragraph(style="Normal")
            paragraph.paragraph_format.left_indent = Cm(0.74)
            paragraph.paragraph_format.right_indent = Cm(0.4)
            paragraph.paragraph_format.first_line_indent = None
            add_inline_content(paragraph, line.lstrip("> "), size=11)
            index += 1
            continue
        if re.match(r"^\[\d+\]", line):
            paragraph = doc.add_paragraph(style="Normal")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.left_indent = Cm(0.74)
            paragraph.paragraph_format.first_line_indent = Cm(-0.74)
            paragraph.paragraph_format.line_spacing = 1.25
            add_inline_content(paragraph, line, size=10.5)
            index += 1
            continue

        paragraph_lines = [line]
        index += 1
        while index < len(lines):
            following = lines[index].strip()
            if (
                not following
                or re.match(r"^(#{1,3})\s+", following)
                or following.startswith("```")
                or following == "\\["
                or re.match(r"^\[建议插图", following)
                or re.match(r"^-\s+", following)
                or re.match(r"^\d+\.\s+", following)
                or following.startswith("|")
                or following.startswith(">")
                or re.match(r"^\[\d+\]", following)
            ):
                break
            paragraph_lines.append(following)
            index += 1
        add_body_paragraph(doc, "".join(paragraph_lines))

    if table_index != len(TABLE_CAPTIONS):
        raise RuntimeError(f"expected {len(TABLE_CAPTIONS)} tables, generated {table_index}")


def extract_source_parts() -> tuple[str, str, str, str, list[str], list[tuple[int, str]]]:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    abstract_index = lines.index("# 摘要")
    keywords_index = lines.index("# 关键词")
    english_index = lines.index("# Abstract")
    english_keywords_index = lines.index("# Key words")
    body_index = next(index for index, line in enumerate(lines) if line.startswith("# 1 "))

    chinese_abstract = "".join(line.strip() for line in lines[abstract_index + 1 : keywords_index] if line.strip())
    chinese_keywords = "".join(line.strip() for line in lines[keywords_index + 1 : english_index] if line.strip())
    english_abstract = " ".join(line.strip() for line in lines[english_index + 1 : english_keywords_index] if line.strip())
    english_keywords = " ".join(line.strip() for line in lines[english_keywords_index + 1 : body_index] if line.strip())
    body_lines = lines[body_index:]
    headings: list[tuple[int, str]] = []
    for line in body_lines:
        match = re.match(r"^(#{1,3})\s+(.+)$", line.strip())
        if match:
            headings.append((len(match.group(1)), match.group(2).strip()))
    return chinese_abstract, chinese_keywords, english_abstract, english_keywords, body_lines, headings


def remove_body_after_marker(doc: Document, marker: Paragraph) -> None:
    body = doc._body._element
    marker_element = marker._p
    seen_marker = False
    for child in list(body):
        if child is marker_element:
            seen_marker = True
            continue
        if seen_marker and child.tag != qn("w:sectPr"):
            body.remove(child)


def clone_toc_template(template_element, text: str, doc: Document):
    clone = deepcopy(template_element)
    p_pr = clone.find(qn("w:pPr"))
    for child in list(clone):
        if child is not p_pr:
            clone.remove(child)
    paragraph = Paragraph(clone, doc._body)
    run = paragraph.add_run(text)
    set_run_font(run, size=12)
    return clone


def roman(number: int) -> str:
    values = ((10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I"))
    result = ""
    for value, token in values:
        while number >= value:
            result += token
            number -= value
    return result


def normalized(text: str) -> str:
    return re.sub(r"\s+", "", text).replace("–", "-").replace("—", "-")


def extract_page_map(pdf_path: Path | None, headings: list[tuple[int, str]]) -> dict[str, str]:
    if pdf_path is None or not pdf_path.exists():
        return {}
    reader = PdfReader(str(pdf_path))
    raw_pages = [page.extract_text() or "" for page in reader.pages]
    compact_pages = [normalized(text) for text in raw_pages]

    def find_page(label: str, *, start=1, last=False) -> int | None:
        needles = {normalized(label), normalized(SPECIAL_HEADINGS.get(label, label))}
        page_numbers = range(start, len(compact_pages) + 1)
        if last:
            page_numbers = range(len(compact_pages), start - 1, -1)
        for page_number in page_numbers:
            if any(needle and needle in compact_pages[page_number - 1] for needle in needles):
                return page_number
        return None

    # The first chapter title also appears in the static contents. Its final
    # occurrence is the real main-text start; using the first occurrence would
    # assign contents-page locations to every chapter.
    body_start = find_page(headings[0][1], start=1, last=True)
    if body_start is None:
        raise RuntimeError("could not locate first body heading in rendered PDF")

    result: dict[str, str] = {}
    for front_label in ("摘要", "Abstract", "目录"):
        page = find_page(front_label, start=1)
        if page is not None and page > 1:
            result[front_label] = roman(page - 1)
    search_start = body_start
    for _level, heading in headings:
        page = find_page(heading, start=search_start, last=heading in SPECIAL_HEADINGS)
        if page is None:
            continue
        result[heading] = f"- {page - body_start + 1} -"
        search_start = page
    return result


def replace_toc(doc: Document, headings: list[tuple[int, str]], page_map: dict[str, str]) -> None:
    paragraphs = doc.paragraphs
    title = paragraphs[59]
    section_marker = paragraphs[124]
    templates = {
        1: deepcopy(paragraphs[60]._p),
        2: deepcopy(paragraphs[63]._p),
        3: deepcopy(paragraphs[76]._p),
    }
    body = doc._body._element
    start_index = list(body).index(title._p) + 1
    marker_index = list(body).index(section_marker._p)
    for child in list(body)[start_index:marker_index]:
        body.remove(child)

    toc_entries: list[tuple[int, str, str]] = [
        (1, "摘    要", page_map.get("摘要", "")),
        (1, "Abstract", page_map.get("Abstract", "")),
        (1, "目    录", page_map.get("目录", "")),
    ]
    for level, heading in headings:
        toc_entries.append((level, SPECIAL_HEADINGS.get(heading, heading), page_map.get(heading, "")))

    for level, text, page in toc_entries:
        clone = clone_toc_template(templates[level], f"{text}\t{page}", doc)
        body.insert(body.index(section_marker._p), clone)


def fix_cover_height(doc: Document) -> None:
    for paragraph in doc.paragraphs[13:20]:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = Pt(2)
    date_paragraph = doc.paragraphs[20]
    date_paragraph.paragraph_format.space_before = Pt(145)
    date_paragraph.paragraph_format.space_after = Pt(0)
    date_paragraph.paragraph_format.line_spacing = 1


def remove_manual_page_break(paragraph: Paragraph) -> None:
    for page_break in paragraph._p.xpath('.//w:br[@w:type="page"]'):
        page_break.getparent().remove(page_break)


def update_headers(doc: Document) -> None:
    seen_parts: set[str] = set()
    sample = "论文标题（选择全部内容后查找替换即可完成相关部分的配置）"
    for section in doc.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            partname = str(header.part.partname)
            if partname in seen_parts:
                continue
            seen_parts.add(partname)
            for paragraph in header.paragraphs:
                for node in paragraph._p.xpath(".//w:t"):
                    if node.text and sample in node.text:
                        node.text = node.text.replace(sample, TITLE)


def set_update_fields(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def _zipinfo_clone(info: ZipInfo) -> ZipInfo:
    clone = ZipInfo(info.filename, info.date_time)
    clone.compress_type = ZIP_DEFLATED
    clone.comment = info.comment
    clone.extra = info.extra
    clone.internal_attr = info.internal_attr
    clone.external_attr = info.external_attr
    clone.create_system = info.create_system
    return clone


def clean_comments(path: Path) -> None:
    skip = {
        "word/comments.xml",
        "word/commentsExtended.xml",
        "word/_rels/comments.xml.rels",
    }
    with tempfile.TemporaryDirectory(prefix="thesis-clean-") as temp_dir:
        rewritten = Path(temp_dir) / path.name
        with ZipFile(path, "r") as source, ZipFile(rewritten, "w", compression=ZIP_DEFLATED) as target:
            for info in source.infolist():
                if info.filename in skip:
                    continue
                data = source.read(info.filename)
                if info.filename == "word/document.xml":
                    root = etree.fromstring(data)
                    for tag in ("commentRangeStart", "commentRangeEnd", "commentReference"):
                        for node in root.xpath(f"//w:{tag}", namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}):
                            node.getparent().remove(node)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
                elif info.filename == "word/_rels/document.xml.rels":
                    root = etree.fromstring(data)
                    for relationship in list(root):
                        rel_type = relationship.get("Type", "")
                        if "comments" in rel_type:
                            root.remove(relationship)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
                elif info.filename == "[Content_Types].xml":
                    root = etree.fromstring(data)
                    for override in list(root):
                        part_name = override.get("PartName", "")
                        if part_name in ("/word/comments.xml", "/word/commentsExtended.xml"):
                            root.remove(override)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
                target.writestr(_zipinfo_clone(info), data)
        shutil.copyfile(rewritten, path)


def build(template: Path, output: Path, toc_pdf: Path | None) -> None:
    if not template.exists():
        raise FileNotFoundError(template)
    actual_hash = sha256(template)
    if actual_hash != EXPECTED_TEMPLATE_SHA256:
        raise RuntimeError(f"template SHA-256 mismatch: {actual_hash}")

    chinese_abstract, chinese_keywords, english_abstract, english_keywords, body_lines, headings = extract_source_parts()
    page_map = extract_page_map(toc_pdf, headings)

    doc = Document(str(template))
    if len(doc.sections) != 5 or len(doc.paragraphs) < 329:
        raise RuntimeError("template structure no longer matches distilled contract")
    configure_body_styles(doc)
    original_paragraphs = doc.paragraphs

    set_multiline_title(original_paragraphs[8])
    replace_cover_identity(original_paragraphs[12])
    expand_cover_identity_box(original_paragraphs[12])
    set_paragraph_text(original_paragraphs[20], "2027年5月", size=16)
    original_paragraphs[20].alignment = WD_ALIGN_PARAGRAPH.CENTER
    fix_cover_height(doc)

    replace_abstract_metadata(original_paragraphs[21])
    set_paragraph_text(original_paragraphs[28], chinese_abstract, size=12)
    original_paragraphs[28].alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    original_paragraphs[28].paragraph_format.first_line_indent = Pt(24)
    original_paragraphs[28].paragraph_format.line_spacing = 1.25
    for index in range(29, 36):
        clear_paragraph(original_paragraphs[index])
    set_paragraph_text(original_paragraphs[36], f"关键词：{chinese_keywords}", size=12)
    original_paragraphs[36].style = doc.styles["Normal"]
    original_paragraphs[36].paragraph_format.first_line_indent = None
    original_paragraphs[36].paragraph_format.line_spacing = 1.25

    replace_english_metadata(original_paragraphs[40])
    remove_manual_page_break(original_paragraphs[39])
    set_paragraph_text(original_paragraphs[43], english_abstract, east=FONT_LATIN, size=12)
    original_paragraphs[43].alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    original_paragraphs[43].paragraph_format.first_line_indent = Pt(24)
    original_paragraphs[43].paragraph_format.line_spacing = 1.25
    set_paragraph_text(original_paragraphs[45], f"Key words: {english_keywords}", east=FONT_LATIN, size=12)
    original_paragraphs[45].paragraph_format.first_line_indent = None
    original_paragraphs[45].paragraph_format.line_spacing = 1.25

    replace_toc(doc, headings, page_map)
    main_marker = doc.paragraphs[-1] if len(doc.paragraphs) == 125 else next(
        paragraph for paragraph in doc.paragraphs if paragraph._p.xpath("./w:pPr/w:sectPr") and paragraph.text == ""
    )
    # The preserved contents section-break paragraph is the final direct paragraph after TOC replacement.
    section_breaks = [paragraph for paragraph in doc.paragraphs if paragraph._p.xpath("./w:pPr/w:sectPr")]
    if len(section_breaks) < 4:
        raise RuntimeError("template section-break structure was not preserved")
    main_marker = section_breaks[3]
    remove_body_after_marker(doc, main_marker)

    figure_generator = load_figure_generator()
    generated = figure_generator.generate_figures()
    parse_body(doc, body_lines, generated)

    update_headers(doc)
    set_update_fields(doc)
    properties = doc.core_properties
    properties.title = TITLE
    properties.subject = "本科毕业设计论文"
    properties.author = "柯佳伟"
    properties.keywords = "实时协同编辑; WebSocket; Lamport逻辑时钟; 字段级LWW; Spring Boot"
    properties.comments = ""
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    clean_comments(output)

    if sha256(template) != EXPECTED_TEMPLATE_SHA256:
        raise RuntimeError("retained template changed during generation")
    print(f"saved={output}")
    print(f"bytes={output.stat().st_size}")
    print(f"sha256={sha256(output)}")
    print(f"headings={len(headings)}")
    print(f"toc_pages={'yes' if page_map else 'no'}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--toc-pdf", type=Path)
    args = parser.parse_args()
    build(
        args.template.resolve(),
        args.output.resolve(),
        args.toc_pdf.resolve() if args.toc_pdf else None,
    )


if __name__ == "__main__":
    main()
