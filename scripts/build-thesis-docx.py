from __future__ import annotations

import argparse
import math
import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "deliverables" / "论文内容底稿.md"
DEFAULT_OUTPUT = ROOT / "deliverables" / "毕业论文初稿-待套学院模板.docx"
QA_DIR = ROOT / "artifacts" / "docx-qa"
FIG_DIR = QA_DIR / "figures"

FONT_BODY = "宋体"
FONT_HEADING = "黑体"
FONT_LATIN = "Times New Roman"
FONT_CODE = "Consolas"
FONT_MATH = "Cambria Math"
FONT_UI = Path("C:/Windows/Fonts/msyh.ttc")
FONT_UI_BOLD = Path("C:/Windows/Fonts/msyhbd.ttc")

INK = "1C2529"
MUTED = "5D6B70"
GREEN = "247A62"
GREEN_DARK = "185845"
BLUE = "2E6F9E"
ORANGE = "C77828"
PALE_GREEN = "E8F3EE"
PALE_BLUE = "EAF2F8"
PALE_ORANGE = "FAF0E3"
LIGHT = "F5F7F6"
GRID = "D9DEDC"
WHITE = "FFFFFF"

TOC_HEADINGS = ["摘要", "1 绪论", "2 需求分析", "3 总体设计", "4 系统实现", "5 系统测试", "结论", "致谢", "参考文献"]

FIGURES = {
    "1-1": ("研究技术路线", "route"),
    "3-1": ("系统总体架构", "architecture"),
    "3-2": ("系统功能模块", "modules"),
    "3-3": ("LWW 字段合并流程", "lww"),
    "3-4": ("WebSocket 连接状态", "ws_state"),
    "3-5": ("断线恢复时序", "reconnect"),
    "3-6": ("操作提交时序", "commit"),
    "3-7": ("数据库逻辑 ER 关系（当前未声明 SQL FOREIGN KEY）", "er"),
    "4-1": ("协同文档双端同步界面", "03-document-owner-editor-sync.png"),
    "4-2": ("协同白板卡片同步界面", "06-board-synced.png"),
    "4-3": ("协同看板任务卡跨列后的同步结果", "07-kanban-drag-synced.png"),
    "4-4": ("工作台空间与成员管理界面", "02-owner-members.png"),
    "4-5": ("版本历史恢复结果", "04-history-restored.png"),
    "4-6": ("正文搜索与协同统计界面", "09-search-result.png"),
    "5-1": ("Playwright 三角色权限验证", ("02-owner-members.png", "05-viewer-document-readonly.png")),
    "5-2": ("WebSocket 中断 30003 ms 后自动重连并补齐", "08-editor-reconnected-after-30s.png"),
    "5-3": ("不同客户端规模下的操作确认延迟", "latency"),
    "5-4": ("不同客户端规模下的吞吐量", "throughput"),
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=GRID, size="6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def hide_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_table_header(row) -> None:
    repeat_table_header(row)
    prevent_row_split(row)


def set_run_font(run, east=FONT_BODY, latin=FONT_LATIN, size=12, bold=None, color=INK, italic=None) -> None:
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    run._element.rPr.rFonts.set(qn("w:ascii"), latin)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), latin)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor.from_string(INK)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(0)
    pf.widow_control = True

    for level, size, before, after in ((1, 16, 0, 14), (2, 14, 14, 7), (3, 12, 10, 5)):
        style = styles[f"Heading {level}"]
        style.font.name = FONT_HEADING
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_HEADING)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True
        if level == 1:
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        style.font.size = Pt(12)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        style.paragraph_format.space_after = Pt(0)

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = FONT_CODE
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    code.font.size = Pt(9.5)
    code.font.color.rgb = RGBColor.from_string(INK)
    code.paragraph_format.left_indent = Cm(0.5)
    code.paragraph_format.right_indent = Cm(0.5)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(4)
    code.paragraph_format.line_spacing = 1.15

    if "Caption Custom" not in styles:
        caption = styles.add_style("Caption Custom", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Caption Custom"]
    caption.font.name = FONT_LATIN
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    caption.font.size = Pt(10.5)
    caption.font.color.rgb = RGBColor.from_string(INK)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, size=9, color=MUTED)


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.6)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("基于SpringBoot与WebSocket的多人实时协同编辑系统的设计与实现")
    set_run_font(run, size=9, color=MUTED)
    bottom = OxmlElement("w:pBdr")
    border = OxmlElement("w:bottom")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "4")
    border.set(qn("w:space"), "4")
    border.set(qn("w:color"), GRID)
    bottom.append(border)
    p._p.get_or_add_pPr().append(bottom)

    footer = section.footer
    add_page_field(footer.paragraphs[0])


def add_cover(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("吉林农业科技学院")
    set_run_font(r, east=FONT_HEADING, size=22, bold=True, color="000000")
    p.paragraph_format.space_after = Pt(24)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("本科毕业设计论文初稿")
    set_run_font(r, east=FONT_HEADING, size=18, bold=True, color="000000")
    p.paragraph_format.space_after = Pt(48)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("基于SpringBoot与WebSocket的\n多人实时协同编辑系统的设计与实现")
    set_run_font(r, east=FONT_HEADING, size=22, bold=True, color="000000")
    p.paragraph_format.line_spacing = 1.35

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("待套学院正式论文模板")
    set_run_font(r, east="微软雅黑", size=11, color=MUTED)
    p.paragraph_format.space_after = Pt(58)

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(4.0)
    table.columns[1].width = Cm(8.0)
    hide_table_borders(table)
    rows = [
        ("学生姓名", "柯佳伟"),
        ("学号", "20235102034"),
        ("专业班级", "软件工程一班"),
        ("指导教师", "揣小龙 讲师"),
    ]
    for row, values in zip(table.rows, rows):
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            set_cell_margins(cell, top=100, bottom=100)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = para.add_run(value)
            set_run_font(run, east=FONT_HEADING if idx == 0 else FONT_BODY, size=13, bold=(idx == 0), color="000000")

    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("二〇二六年九月")
    set_run_font(r, size=12, color="000000")
    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading(text, level=level)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_inline_runs(paragraph, text: str, base_size=12) -> None:
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=base_size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, east="微软雅黑", latin=FONT_CODE, size=max(9.5, base_size - 1), color=GREEN_DARK)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=base_size)


def add_body_paragraph(
    doc: Document,
    text: str,
    first_indent=True,
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    font_size=12,
    line_spacing=1.5,
) -> None:
    p = doc.add_paragraph()
    p.alignment = alignment
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.widow_control = True
    if first_indent:
        p.paragraph_format.first_line_indent = Pt(24)
    add_inline_runs(p, text, base_size=font_size)


def add_abstract(doc: Document, abstract_lines: list[str], keyword_lines: list[str]) -> None:
    add_heading(doc, "摘要", 1)
    for line in abstract_lines:
        if line.strip() and not line.startswith(">"):
            add_body_paragraph(doc, line.strip(), font_size=11, line_spacing=1.35)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    r = p.add_run("关键词：")
    set_run_font(r, east=FONT_HEADING, size=12, bold=True, color="000000")
    keywords = "".join(line.strip() for line in keyword_lines if line.strip())
    r = p.add_run(keywords)
    set_run_font(r, size=12, color="000000")


def extract_toc_pages(pdf_path: Path | None) -> dict[str, int]:
    if not pdf_path or not pdf_path.exists():
        return {}
    reader = PdfReader(str(pdf_path))
    raw_texts = [page.extract_text() or "" for page in reader.pages]
    texts = [re.sub(r"\s+", "", text) for text in raw_texts]

    def appears_as_heading(heading: str, page_index: int) -> bool:
        key = re.sub(r"\s+", "", heading)
        lines = [re.sub(r"\s+", "", line) for line in raw_texts[page_index].splitlines()[:10]]
        if key in lines:
            return True
        if key and key[0].isdigit():
            for line in lines:
                if line.startswith(key[0]) and all(character in line for character in key):
                    return True
        return False

    pages = {}
    for page_number in range(1, len(texts) + 1):
        if appears_as_heading(TOC_HEADINGS[0], page_number - 1):
            pages[TOC_HEADINGS[0]] = page_number
            break

    toc_page = 0
    for page_number, text in enumerate(texts, start=1):
        if "目录" in text and "1绪论" in text and "2需求分析" in text:
            toc_page = page_number
            break

    cursor = toc_page + 1
    for heading in TOC_HEADINGS[1:]:
        for page_number in range(cursor, len(texts) + 1):
            if appears_as_heading(heading, page_number - 1):
                pages[heading] = page_number
                cursor = page_number + 1
                break
    return pages


def add_toc(doc: Document, page_map: dict[str, int]) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("目录")
    set_run_font(r, east=FONT_HEADING, size=18, bold=True, color="000000")
    for heading in TOC_HEADINGS:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(15.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = p.add_run(heading)
        set_run_font(r, east=FONT_HEADING if re.match(r"^\d+\s", heading) else FONT_BODY, size=12, bold=bool(re.match(r"^\d+\s", heading)))
        p.add_run("\t")
        r = p.add_run(str(page_map.get(heading, "")))
        set_run_font(r, size=11)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("目录页码依据本初稿渲染结果生成；套用学院模板后需重新更新。")
    set_run_font(r, size=9.5, color=MUTED)


def add_code_block(doc: Document, lines: list[str]) -> None:
    p = doc.add_paragraph(style="Code Block")
    p.paragraph_format.keep_together = True
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F5F4")
    p._p.get_or_add_pPr().append(shading)
    for idx, line in enumerate(lines):
        run = p.add_run(line)
        set_run_font(run, east="微软雅黑", latin=FONT_CODE, size=9.5, color=INK)
        if idx < len(lines) - 1:
            run.add_break()


def add_equation(doc: Document, lines: list[str]) -> None:
    text = " ".join(line.strip() for line in lines if line.strip() not in ("\\[", "\\]"))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run_font(r, east=FONT_MATH, latin=FONT_MATH, size=12, color="000000")


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_table_separator(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    rows = [row + [""] * (cols - len(row)) for row in rows]
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    font_size = 9 if cols >= 8 else 10 if cols >= 5 else 10.5
    for r_idx, row_values in enumerate(rows):
        row = table.rows[r_idx]
        prevent_row_split(row)
        if r_idx == 0:
            set_repeat_table_header(row)
        for c_idx, value in enumerate(row_values):
            cell = row.cells[c_idx]
            set_cell_border(cell)
            set_cell_margins(cell, top=80, start=85, bottom=80, end=85)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "2F5F67" if r_idx == 0 else ("F3F7F6" if r_idx % 2 == 0 else WHITE))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (r_idx == 0 or len(value) < 24) else WD_ALIGN_PARAGRAPH.LEFT
            clean = value.replace("**", "")
            run = p.add_run(clean)
            set_run_font(run, east="微软雅黑" if r_idx == 0 else FONT_BODY, size=font_size, bold=(r_idx == 0), color=WHITE if r_idx == 0 else INK)
    before = doc.add_paragraph()
    before.paragraph_format.space_after = Pt(0)


def font(size: int, bold=False) -> ImageFont.FreeTypeFont:
    path = FONT_UI_BOLD if bold and FONT_UI_BOLD.exists() else FONT_UI
    return ImageFont.truetype(str(path), size=size)


def draw_centered(draw: ImageDraw.ImageDraw, box, text, size=30, color="#1C2529", bold=False, max_lines=3) -> None:
    x1, y1, x2, y2 = box
    f = font(size, bold)
    max_width = max(1, x2 - x1 - 30)
    lines = []
    for source_line in text.split("\n"):
        current = ""
        # Keep Latin identifiers intact while still allowing Chinese text to wrap by character.
        tokens = re.findall(r"[A-Za-z0-9_./:+-]+|\s+|.", source_line)
        for token in tokens:
            candidate = current + token
            if draw.textbbox((0, 0), candidate, font=f)[2] <= max_width or not current:
                current = candidate
            else:
                lines.append(current.rstrip())
                current = token.lstrip()
        if current:
            lines.append(current.rstrip())
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        lines[-1] = lines[-1][:-1] + "…"
    line_h = size + 12
    total_h = line_h * len(lines)
    y = y1 + (y2 - y1 - total_h) / 2
    for line in lines:
        bounds = draw.textbbox((0, 0), line, font=f)
        draw.text((x1 + (x2 - x1 - (bounds[2] - bounds[0])) / 2, y), line, font=f, fill=color)
        y += line_h


def arrow(draw, start, end, color="#2E6F9E", width=6) -> None:
    draw.line([start, end], fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 20
    left = (end[0] - length * math.cos(angle - 0.55), end[1] - length * math.sin(angle - 0.55))
    right = (end[0] - length * math.cos(angle + 0.55), end[1] - length * math.sin(angle + 0.55))
    draw.polygon([end, left, right], fill=color)


def rounded_box(draw, box, text, fill="#E8F3EE", outline="#247A62", size=28, bold=False) -> None:
    draw.rounded_rectangle(box, radius=16, fill=fill, outline=outline, width=4)
    draw_centered(draw, box, text, size=size, bold=bold)


def canvas(title: str, width=1800, height=950):
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.text((70, 35), title, font=font(42, True), fill="#1C2529")
    draw.line((70, 100, width - 70, 100), fill="#D9DEDC", width=3)
    return image, draw


def save_route(path: Path) -> None:
    image, draw = canvas("研究技术路线")
    nodes = ["文献与需求分析", "统一元素模型", "Spring Boot 与\nWebSocket", "Vue 三类应用", "MySQL 日志与快照", "分层自动化验证"]
    y1, y2 = 340, 560
    gap = 28
    box_w = (1660 - gap * 5) / 6
    for idx, label in enumerate(nodes):
        x1 = 70 + idx * (box_w + gap)
        x2 = x1 + box_w
        rounded_box(draw, (x1, y1, x2, y2), label, fill=["#E8F3EE", "#EAF2F8", "#F5F7F6"][idx % 3], size=25, bold=idx in (1, 5))
        if idx < len(nodes) - 1:
            arrow(draw, (x2 + 4, 450), (x2 + gap - 4, 450))
    draw_centered(draw, (200, 650, 1600, 800), "范围护栏：段落级字段 LWW · 结构化卡片 · 单实例协调", size=31, color="#C77828", bold=True)
    image.save(path)


def save_architecture(path: Path) -> None:
    image, draw = canvas("系统总体架构")
    rounded_box(draw, (100, 165, 1700, 265), "浏览器 A / B / C", fill="#F5F7F6", size=32, bold=True)
    rounded_box(draw, (100, 340, 1700, 440), "Vue 3 + Vite 工作台与三类编辑器", fill="#E8F3EE", size=32, bold=True)
    rounded_box(draw, (100, 520, 820, 655), "REST 控制面\n认证 · 空间 · 历史 · 检索 · 导出 · 统计", fill="#EAF2F8", outline="#2E6F9E", size=27)
    rounded_box(draw, (980, 520, 1700, 655), "WebSocket 数据面\nhello/op · incremental/applied · presence", fill="#EAF2F8", outline="#2E6F9E", size=27)
    rounded_box(draw, (100, 750, 820, 870), "自研 engine\nElement + Op + Lamport + LWW + 墓碑", fill="#E8F3EE", size=27, bold=True)
    rounded_box(draw, (980, 750, 1700, 870), "MySQL 8\n业务表 + operation_log + element_snapshot", fill="#FAF0E3", outline="#C77828", size=27, bold=True)
    arrow(draw, (900, 265), (900, 330))
    arrow(draw, (470, 440), (470, 510))
    arrow(draw, (1330, 440), (1330, 510))
    arrow(draw, (470, 655), (470, 740))
    arrow(draw, (1330, 655), (1330, 740), color="#C77828")
    draw.text((1140, 690), "先持久化，后广播", font=font(25, True), fill="#C77828")
    image.save(path)


def save_modules(path: Path) -> None:
    image, draw = canvas("系统功能模块")
    labels = [
        ("认证与权限", "JWT · BCrypt · owner/editor/viewer"),
        ("空间与内容", "空间 · 成员 · 文档元数据"),
        ("实时协同", "房间 · 串行提交 · 在线状态"),
        ("三类编辑器", "文档 · 白板 · 看板"),
        ("产品能力", "历史 · 检索 · 导出 · 统计"),
        ("持久化与测试", "日志 · 快照 · JUnit · Playwright"),
    ]
    for idx, (title, sub) in enumerate(labels):
        row, col = divmod(idx, 3)
        x1 = 90 + col * 570
        y1 = 175 + row * 340
        box = (x1, y1, x1 + 500, y1 + 245)
        draw.rounded_rectangle(
            box,
            radius=16,
            fill="#E8F3EE" if row == 0 else "#EAF2F8",
            outline="#247A62",
            width=4,
        )
        draw_centered(draw, (x1 + 18, y1 + 30, x1 + 482, y1 + 102), title, size=30, bold=True)
        draw_centered(draw, (x1 + 18, y1 + 112, x1 + 482, y1 + 218), sub, size=23, max_lines=2)
    draw_centered(draw, (250, 800, 1550, 900), "公共约束：统一 Element + Op 协议与空间级资源权限", size=30, color="#185845", bold=True)
    image.save(path)


def save_lww(path: Path) -> None:
    image, draw = canvas("LWW 字段合并流程")
    nodes = [
        ((100, 250, 380, 400), "接收操作"),
        ((520, 250, 800, 400), "校验类型、字段\n和时钟"),
        ((940, 250, 1220, 400), "比较 writtenAt\n与 op.clock"),
        ((1420, 170, 1710, 320), "新时钟更大\n写入字段"),
        ((1420, 450, 1710, 600), "旧时钟或相等\n忽略操作"),
    ]
    for idx, (box, label) in enumerate(nodes):
        rounded_box(draw, box, label, fill="#EAF2F8" if idx < 3 else ("#E8F3EE" if idx == 3 else "#F5F7F6"), outline="#2E6F9E" if idx < 3 else "#247A62", size=28, bold=True)
    arrow(draw, (380, 325), (510, 325))
    arrow(draw, (800, 325), (930, 325))
    arrow(draw, (1220, 300), (1410, 245))
    arrow(draw, (1220, 350), (1410, 525), color="#5D6B70")
    draw.text((1255, 220), "是", font=font(24, True), fill="#247A62")
    draw.text((1260, 430), "否", font=font(24, True), fill="#5D6B70")
    rounded_box(draw, (420, 700, 1380, 835), "DELETE 写入 #tombstone=true，并参加同一套字段级时钟裁决", fill="#FAF0E3", outline="#C77828", size=28, bold=True)
    image.save(path)


def save_ws_state(path: Path) -> None:
    image, draw = canvas("WebSocket 连接状态机", height=820)
    labels = ["未连接", "连接中", "同步中", "在线", "退避重连"]
    node_y = 320
    centers = [(200, node_y), (540, node_y), (880, node_y), (1220, node_y), (1560, node_y)]
    for idx, (label, (cx, cy)) in enumerate(zip(labels, centers)):
        fill = "#E8F3EE" if label == "在线" else "#EAF2F8"
        rounded_box(draw, (cx - 125, cy - 80, cx + 125, cy + 80), label, fill=fill, size=30, bold=True)
        if idx < len(labels) - 1:
            arrow(draw, (cx + 130, cy), (centers[idx + 1][0] - 135, cy))
    loop_y = 535
    arrow(draw, (1560, node_y + 110), (1200, loop_y), color="#C77828")
    arrow(draw, (1200, loop_y), (540, node_y + 110), color="#C77828")
    draw_centered(
        draw,
        (430, loop_y + 20, 1370, 760),
        "断线保留 lastRev；可连续则补增量，\n否则回退到全量快照",
        size=27,
        color="#C77828",
        bold=True,
    )
    image.save(path)


def save_reconnect(path: Path) -> None:
    image, draw = canvas("断线 30 秒与自动恢复")
    x_client, x_server = 420, 1380
    draw.line((x_client, 180, x_client, 870), fill="#247A62", width=5)
    draw.line((x_server, 180, x_server, 870), fill="#2E6F9E", width=5)
    draw_centered(draw, (250, 120, 590, 190), "编辑者浏览器", size=30, bold=True)
    draw_centered(draw, (1210, 120, 1550, 190), "协同服务", size=30, bold=True)
    events = [
        (270, "hello(lastRev=20)", "right", "#2E6F9E"),
        (390, "welcome / incremental", "left", "#247A62"),
        (525, "WebSocket 传输中断 30003 ms", "center", "#C77828"),
        (700, "hello(lastRev=20)", "right", "#2E6F9E"),
        (815, "incremental rev 21..24", "left", "#247A62"),
    ]
    for y, label, direction, color in events:
        if direction == "right":
            arrow(draw, (x_client + 15, y), (x_server - 15, y), color=color)
            draw_centered(draw, (600, y - 64, 1200, y - 12), label, size=25, color=color, bold=True)
        elif direction == "left":
            arrow(draw, (x_server - 15, y), (x_client + 15, y), color=color)
            draw_centered(draw, (600, y - 64, 1200, y - 12), label, size=25, color=color, bold=True)
        else:
            draw.rounded_rectangle((560, y - 48, 1240, y + 48), radius=16, fill="#FAF0E3", outline=color, width=4)
            draw_centered(draw, (600, y - 45, 1200, y + 45), label, size=25, color=color, bold=True)
    image.save(path)


def save_commit(path: Path) -> None:
    image, draw = canvas("操作提交时序：数据库成功后才广播")
    labels = ["Vue 客户端", "Handler", "CollabRoom", "OpStore / MySQL", "engine", "其他客户端"]
    xs = [120, 420, 720, 1040, 1340, 1640]
    for x, label in zip(xs, labels):
        draw.line((x, 210, x, 865), fill="#B9C4C1", width=3)
        draw_centered(draw, (x - 115, 125, x + 115, 205), label, size=24, bold=True)
    steps = [
        (280, 0, 1, "op"),
        (375, 1, 2, "校验后入队"),
        (470, 2, 3, "追加 operation_log"),
        (565, 3, 2, "持久化成功"),
        (660, 2, 4, "apply + rev"),
        (755, 2, 5, "broadcast applied"),
    ]
    for y, a, b, label in steps:
        color = "#C77828" if "持久化" in label or "operation" in label else "#2E6F9E"
        start = (xs[a] + (10 if b > a else -10), y)
        end = (xs[b] - (10 if b > a else -10), y)
        arrow(draw, start, end, color=color)
        draw_centered(draw, (min(xs[a], xs[b]) + 15, y - 48, max(xs[a], xs[b]) - 15, y - 5), label, size=22, color=color, bold=True)
    image.save(path)


def save_er(path: Path) -> None:
    image, draw = canvas("数据库逻辑实体关系")
    entities = {
        "user": (100, 180, 520, 350),
        "space": (690, 180, 1110, 350),
        "space_member": (1280, 180, 1700, 350),
        "document": (100, 560, 520, 730),
        "operation_log": (690, 560, 1110, 730),
        "element_snapshot": (1280, 560, 1700, 730),
    }
    details = {
        "user": "id · username · password_hash",
        "space": "id · name · owner_id",
        "space_member": "space_id · user_id · role",
        "document": "doc_key · space_id · kind",
        "operation_log": "doc_id · seq · operation_id",
        "element_snapshot": "doc_id · rev · state_json",
    }
    for name, box in entities.items():
        rounded_box(draw, box, f"{name}\n{details[name]}", fill="#EAF2F8" if "log" not in name and "snapshot" not in name else "#E8F3EE", size=25, bold=True)
    arrow(draw, (520, 265), (680, 265))
    arrow(draw, (1110, 265), (1270, 265))
    arrow(draw, (310, 350), (310, 550))
    arrow(draw, (520, 650), (680, 650))
    arrow(draw, (1110, 650), (1270, 650))
    draw_centered(draw, (200, 800, 1600, 900), "图中为逻辑关联；当前 DDL 未声明 SQL FOREIGN KEY", size=28, color="#C77828", bold=True)
    image.save(path)


def save_latency(path: Path) -> None:
    image, draw = canvas("并发客户端与操作确认延迟（ms）")
    left, top, right, bottom = 180, 180, 1690, 790
    draw.line((left, top, left, bottom), fill="#5D6B70", width=4)
    draw.line((left, bottom, right, bottom), fill="#5D6B70", width=4)
    values = {"P50": [16.31, 46.30, 85.14], "P95": [31.04, 73.88, 115.37], "P99": [34.00, 98.28, 129.71]}
    colors = {"P50": "#247A62", "P95": "#2E6F9E", "P99": "#C77828"}
    max_y = 220
    for tick in range(0, 221, 50):
        y = bottom - (tick / max_y) * (bottom - top)
        draw.line((left, y, right, y), fill="#E4E9E7", width=2)
        draw.text((100, y - 16), str(tick), font=font(22), fill="#5D6B70")
    threshold_y = bottom - (200 / max_y) * (bottom - top)
    draw.line((left, threshold_y, right, threshold_y), fill="#C77828", width=4)
    draw.text((right - 290, threshold_y - 38), "50 客户端 P95 门限 200 ms", font=font(22, True), fill="#C77828")
    clients = [10, 50, 100]
    group_w = 350
    bar_w = 72
    for g, client in enumerate(clients):
        center = 420 + g * 500
        for idx, name in enumerate(("P50", "P95", "P99")):
            value = values[name][g]
            x1 = center - 120 + idx * 95
            y1 = bottom - (value / max_y) * (bottom - top)
            draw.rectangle((x1, y1, x1 + bar_w, bottom), fill=colors[name])
            draw_centered(draw, (x1 - 8, y1 - 48, x1 + bar_w + 8, y1 - 5), f"{value:.2f}", size=18, color=colors[name], bold=True)
        draw_centered(draw, (center - 150, bottom + 12, center + 150, bottom + 75), f"{client} 客户端", size=24, bold=True)
    legend_y = 880
    for idx, name in enumerate(("P50", "P95", "P99")):
        x = 650 + idx * 220
        draw.rectangle((x, legend_y, x + 34, legend_y + 34), fill=colors[name])
        draw.text((x + 48, legend_y - 1), name, font=font(24, True), fill="#1C2529")
    image.save(path)


def save_throughput(path: Path) -> None:
    image, draw = canvas("同房间 closed-loop 吞吐量（ops/s）")
    left, top, right, bottom = 220, 180, 1640, 790
    draw.line((left, top, left, bottom), fill="#5D6B70", width=4)
    draw.line((left, bottom, right, bottom), fill="#5D6B70", width=4)
    values = [473, 988, 1151]
    clients = [10, 50, 100]
    max_y = 1300
    for tick in range(0, 1301, 250):
        y = bottom - (tick / max_y) * (bottom - top)
        draw.line((left, y, right, y), fill="#E4E9E7", width=2)
        draw.text((120, y - 16), str(tick), font=font(22), fill="#5D6B70")
    for idx, (client, value) in enumerate(zip(clients, values)):
        center = 470 + idx * 460
        y = bottom - (value / max_y) * (bottom - top)
        color = ["#247A62", "#2E6F9E", "#C77828"][idx]
        draw.rounded_rectangle((center - 115, y, center + 115, bottom), radius=12, fill=color)
        draw_centered(draw, (center - 150, y - 65, center + 150, y - 10), f"{value}", size=28, color=color, bold=True)
        draw_centered(draw, (center - 150, bottom + 15, center + 150, bottom + 75), f"{client} 客户端", size=25, bold=True)
    draw_centered(draw, (300, 850, 1500, 925), "每客户端 20 ops；三档 lost=0、duplicate=0", size=28, color="#185845", bold=True)
    image.save(path)


def generate_figures() -> dict[str, Path]:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    generators = {
        "route": save_route,
        "architecture": save_architecture,
        "modules": save_modules,
        "lww": save_lww,
        "ws_state": save_ws_state,
        "reconnect": save_reconnect,
        "commit": save_commit,
        "er": save_er,
        "latency": save_latency,
        "throughput": save_throughput,
    }
    paths = {}
    for key, generator in generators.items():
        path = FIG_DIR / f"{key}.png"
        generator(path)
        paths[key] = path
    return paths


def add_caption(doc: Document, number: str, title: str) -> None:
    p = doc.add_paragraph(style="Caption Custom")
    p.add_run(f"图{number} {title}")


def add_image(doc: Document, path: Path, number: str, title: str, max_width_cm=15.3, max_height_cm=10.8) -> None:
    with Image.open(path) as image:
        width_px, height_px = image.size
    aspect = width_px / height_px
    width_cm = max_width_cm
    height_cm = width_cm / aspect
    if height_cm > max_height_cm:
        height_cm = max_height_cm
        width_cm = height_cm * aspect
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm), height=Cm(height_cm))
    add_caption(doc, number, title)


def add_two_images(doc: Document, paths: tuple[Path, Path], number: str, title: str) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for cell, path in zip(table.rows[0].cells, paths):
        cell.width = Cm(7.5)
        set_cell_border(cell, color=GRID, size="4")
        set_cell_margins(cell, top=60, start=60, bottom=60, end=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        with Image.open(path) as image:
            aspect = image.width / image.height
        width_cm = 7.1
        height_cm = min(5.2, width_cm / aspect)
        p.add_run().add_picture(str(path), width=Cm(width_cm), height=Cm(height_cm))
    add_caption(doc, number, title)


def handle_figure(doc: Document, number: str, generated: dict[str, Path]) -> None:
    if number not in FIGURES:
        return
    title, source = FIGURES[number]
    screenshot_dir = ROOT / "artifacts" / "e2e-vue-final"
    if isinstance(source, tuple):
        add_two_images(doc, (screenshot_dir / source[0], screenshot_dir / source[1]), number, title)
    elif source.endswith(".png"):
        add_image(doc, screenshot_dir / source, number, title)
    else:
        add_image(doc, generated[source], number, title, max_height_cm=9.2)


def parse_body(doc: Document, lines: list[str], generated: dict[str, Path]) -> None:
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1:
                doc.add_page_break()
            add_heading(doc, text, level)
            i += 1
            continue
        figure = re.match(r"^\[建议插图(\d+-\d+)：.*\]$", line)
        if figure:
            handle_figure(doc, figure.group(1), generated)
            i += 1
            continue
        if line.startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i].rstrip())
                i += 1
            add_code_block(doc, code_lines)
            i += 1
            continue
        if line == "\\[":
            equation_lines = [line]
            i += 1
            while i < len(lines):
                equation_lines.append(lines[i])
                if lines[i].strip() == "\\]":
                    i += 1
                    break
                i += 1
            add_equation(doc, equation_lines)
            continue
        if line.startswith("|") and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            rows = [split_table_row(line)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_table_row(lines[i]))
                i += 1
            add_markdown_table(doc, rows)
            continue
        if re.match(r"^-\s+", line):
            while i < len(lines) and re.match(r"^-\s+", lines[i].strip()):
                text = re.sub(r"^-\s+", "", lines[i].strip())
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.left_indent = Cm(0.74)
                p.paragraph_format.first_line_indent = Cm(-0.32)
                add_inline_runs(p, text)
                i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i].strip()):
                text = re.sub(r"^\d+\.\s+", "", lines[i].strip())
                p = doc.add_paragraph(style="List Number")
                p.paragraph_format.left_indent = Cm(0.74)
                p.paragraph_format.first_line_indent = Cm(-0.32)
                add_inline_runs(p, text)
                i += 1
            continue
        if line.startswith(">"):
            text = line.lstrip("> ")
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.right_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            add_inline_runs(p, text, base_size=11)
            i += 1
            continue
        if re.match(r"^\[\d+\]", line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.74)
            p.paragraph_format.first_line_indent = Cm(-0.74)
            p.paragraph_format.line_spacing = 1.3
            add_inline_runs(p, line, base_size=10.5)
            i += 1
            continue

        paragraph_lines = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or re.match(r"^(#{1,3})\s+", nxt) or nxt.startswith("```") or nxt == "\\[" or re.match(r"^\[建议插图", nxt) or re.match(r"^-\s+", nxt) or re.match(r"^\d+\.\s+", nxt) or nxt.startswith("|") or nxt.startswith(">") or re.match(r"^\[\d+\]", nxt):
                break
            paragraph_lines.append(nxt)
            i += 1
        add_body_paragraph(doc, "".join(paragraph_lines))


def build(output: Path, toc_pdf: Path | None) -> None:
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    abstract_idx = lines.index("# 摘要")
    keywords_idx = lines.index("# 关键词")
    body_idx = next(idx for idx, line in enumerate(lines) if line.startswith("# 1 "))
    abstract_lines = lines[abstract_idx + 1:keywords_idx]
    keyword_lines = lines[keywords_idx + 1:body_idx]
    body_lines = lines[body_idx:]

    generated = generate_figures()
    doc = Document()
    configure_styles(doc)
    setup_document(doc)
    props = doc.core_properties
    props.title = "基于SpringBoot与WebSocket的多人实时协同编辑系统的设计与实现"
    props.subject = "本科毕业设计论文初稿"
    props.author = "柯佳伟"
    props.keywords = "实时协同编辑, WebSocket, Lamport逻辑时钟, Last-Write-Wins, Spring Boot, Vue 3"
    props.comments = "依据定稿任务书、开题报告和当前实现生成；待套学院正式模板。"

    add_cover(doc)
    add_abstract(doc, abstract_lines, keyword_lines)
    doc.add_page_break()
    add_toc(doc, extract_toc_pages(toc_pdf))
    parse_body(doc, body_lines, generated)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(f"saved={output}")
    print(f"bytes={output.stat().st_size}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--toc-pdf", type=Path)
    args = parser.parse_args()
    build(args.output.resolve(), args.toc_pdf.resolve() if args.toc_pdf else None)


if __name__ == "__main__":
    main()
